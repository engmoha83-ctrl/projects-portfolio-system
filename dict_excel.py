# -*- coding: utf-8 -*-
"""
القواميس ذهابًا وإيابًا عبر إكسل.

يُنزَّل قالبٌ فيه القواميس الحالية — ورقة لكلّ قاموس، وقوائم منسدلة للحقول
التي تشير إلى قاموس آخر — فيُملأ ويُرفع. والرفع يُعرض فرقُه أوّلًا ولا يُكتب
شيء قبل التأكيد.

**الدمج هو الافتراض.** الصفّ الموجود في الملفّ يُضاف أو يُحدَّث، وما ليس في
الملفّ يبقى كما هو؛ فمن رفع ورقة معدّلات لنوعَي عمل لا يمحو معدّلات الباقي.
والاستبدال خيارٌ صريح: الورقة الموجودة في الملفّ تحلّ محلّ قاموسها كلّه.

الأسماء في الملفّ بالإنجليزية لأنّ الواجهة إنجليزية، والعناوين تُطابَق بعد
تطبيعها فترتيب الأعمدة حرّ.
"""

import io
import re

REL = ("FS", "SS", "FF", "SF")
SCOPES = ("same", "next", "prev", "parent", "children")

# ── تعريف الأوراق: (المفتاح في الملفّ، العنوان، الإلزامي، العرض، القائمة) ──
SHEETS = {
    "Work types": [
        ("code", "Code", True, 10, None), ("name", "Name", True, 30, None),
        ("discipline", "Discipline", False, 12, "disc"),
        ("csi_division", "CSI division", False, 12, "csi"),
        ("unit", "Unit", False, 8, None), ("days", "Default days", False, 12, None),
        ("applies_to", "Applies to", False, 22, None),
        ("keywords", "Keywords", False, 50, None),
        ("per_item", "Activities", False, 14, "gran")],
    "Steps": [
        ("work_type", "Work type", True, 12, "wt"), ("code", "Step code", True, 11, None),
        ("name", "Step name", False, 26, None), ("weight", "Weight %", False, 10, None),
        ("rel_type", "Link to previous", False, 15, "rel"), ("lag_days", "Lag (days)", False, 10, None)],
    "Productivity": [
        ("work_type", "Work type", True, 12, "wt"), ("step", "Step", False, 10, None),
        ("resource", "Resource", True, 12, "res"), ("hours_per_unit", "Hours per unit", False, 14, None),
        ("crew_count", "Crew", False, 8, None), ("driving", "Driving", False, 9, "yn")],
    "Work-type materials": [
        ("work_type", "Work type", True, 12, "wt"), ("material", "Material", True, 12, "mat"),
        ("qty_per_unit", "Qty per unit", False, 13, None)],
    "Logic rules": [
        ("pred_wt", "Predecessor", True, 13, "wt"), ("succ_wt", "Successor", True, 13, "wt"),
        ("scope", "Scope", False, 11, "scope"), ("rel_type", "Type", False, 8, "rel"),
        ("lag_days", "Lag (days)", False, 10, None)],
    "Resources": [
        ("code", "Code", True, 10, None), ("name", "Name", True, 26, None),
        ("kind", "Kind", False, 12, "kind"), ("unit_cost", "Cost per hour", False, 14, None)],
    "Materials": [
        ("code", "Code", True, 10, None), ("name", "Name", True, 26, None),
        ("unit", "Unit", False, 8, None), ("unit_cost", "Unit cost", False, 12, None),
        ("lead_time_days", "Lead time (days)", False, 16, None),
        ("keywords", "Keywords", False, 40, None)],
    "Engineering": [
        ("work_type", "Work type", True, 12, "wt"), ("code", "Code", True, 10, None),
        ("name", "Name", True, 28, None), ("duration_days", "Duration (days)", False, 15, None),
        ("rel_type", "Link to previous", False, 15, "rel"), ("lag_days", "Lag (days)", False, 10, None)],
    "T&C": [
        ("work_type", "Work type", True, 12, "wt"), ("code", "Code", True, 10, None),
        ("name", "Name", True, 28, None), ("duration_days", "Duration (days)", False, 15, None),
        ("level_key", "Repeats at", False, 16, "lvl")],
    "Disciplines": [
        ("code", "Code", True, 10, None), ("name", "Name", True, 30, None),
        ("csi_division", "CSI division", False, 12, "csi")],
}

README = [
    ("How to use this file", ""),
    ("", ""),
    ("1", "Each sheet is one dictionary. Edit rows, add rows at the bottom, or delete rows."),
    ("2", "Columns marked * are required. The order of columns does not matter; the header names do."),
    ("3", "Cells with a drop-down point to another dictionary — pick from the list so the codes match."),
    ("4", "Upload the file from the Dictionaries page. You will see what changes before anything is saved."),
    ("", ""),
    ("Merge (default)", "Rows in the file are added or updated. Anything not in the file stays as it is."),
    ("", "For Steps, Productivity and Work-type materials, a work type that appears in the sheet has its rows "
         "replaced by the file's rows; work types that do not appear keep theirs."),
    ("Replace", "A sheet in the file replaces that whole dictionary. Delete a sheet from the file to leave that "
                "dictionary untouched."),
    ("", ""),
    ("Steps", "Weights are percentages of the work type and should add up to 100. "
              "'Link to previous' is the relation from the step above it."),
    ("Work types", "Activities: Grouped makes one activity per work type in each place, carrying all "
                   "its BOQ items; Per item makes one activity for every BOQ item."),
    ("Productivity", "Hours of that resource per unit of work, and how many of it work together (crew). "
                     "Leave Step empty for a rate that covers the whole work type. "
                     "Driving = Yes makes that resource set the duration."),
    ("Logic rules", "Scope: same, next, prev, parent or children — where the successor sits relative to "
                    "the predecessor's place."),
    ("T&C", "Repeats at: a level name (Floor, Building…) or Project for once per project."),
    ("CSI divisions", "Reference only — fixed for every project and not read on upload."),
]


# ═══════════════════════════ التصدير ═══════════════════════════

def build_workbook(d, blank=False):
    """قالب القواميس: الحاليّة إن لم يُطلب فارغًا، وقوائم منسدلة للإشارات."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    navy = PatternFill("solid", fgColor="1F3A5F")
    tint = PatternFill("solid", fgColor="F1F3F7")
    hfont = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110
    for i, (a, b) in enumerate(README, 1):
        ws.cell(i, 1, a).font = Font(bold=True, size=14 if i == 1 else 11)
        c = ws.cell(i, 2, b)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # القوائم في ورقة مخفيّة، فتبقى القائمة المنسدلة صحيحة مهما طالت
    lists = {
        "wt": [w["code"] for w in d["work_types"]],
        "disc": [x["code"] for x in d["disciplines"]],
        "csi": [c["division"] for c in d["csi"]],
        "res": [r["code"] for r in d["resources"]],
        "mat": [m["code"] for m in d["materials"]],
        "rel": list(REL), "scope": list(SCOPES),
        "kind": ["Labour", "Equipment"], "yn": ["Yes", "No"],
        "gran": ["Grouped", "Per item"],
        "lvl": ["Project"] + [l["label"] for l in d["levels"]],
    }
    lw = wb.create_sheet("Lists")
    ref = {}
    for j, (k, vals) in enumerate(lists.items(), 1):
        col = get_column_letter(j)
        lw.cell(1, j, k)
        for i, v in enumerate(vals, 2):
            lw.cell(i, j, v)
        ref[k] = f"Lists!${col}$2:${col}${max(2, len(vals) + 1)}"
    lw.sheet_state = "hidden"

    rows = {} if blank else _rows_for_export(d)
    for title, cols in SHEETS.items():
        s = wb.create_sheet(title)
        for j, (_k, head, req, width, _l) in enumerate(cols, 1):
            c = s.cell(1, j, head + (" *" if req else ""))
            c.font, c.fill = hfont, navy
            c.alignment = Alignment(vertical="center")
            s.column_dimensions[get_column_letter(j)].width = width
        s.row_dimensions[1].height = 22
        s.freeze_panes = "A2"
        for i, r in enumerate(rows.get(title, []), 2):
            for j, (k, *_rest) in enumerate(cols, 1):
                s.cell(i, j, r.get(k))
        last = max(500, len(rows.get(title, [])) + 300)
        for j, (_k, _h, _r, _w, lst) in enumerate(cols, 1):
            if not lst:
                continue
            dv = DataValidation(type="list", formula1=ref[lst], allow_blank=True,
                                showErrorMessage=lst in ("rel", "scope", "kind", "yn", "lvl"))
            col = get_column_letter(j)
            dv.add(f"{col}2:{col}{last}")
            s.add_data_validation(dv)

    s = wb.create_sheet("CSI divisions")
    for j, h in enumerate(("Division", "Name"), 1):
        c = s.cell(1, j, h)
        c.font, c.fill = hfont, navy
    s.column_dimensions["A"].width = 10
    s.column_dimensions["B"].width = 50
    for i, c in enumerate(d["csi"], 2):
        s.cell(i, 1, c["division"])
        s.cell(i, 2, c["name"]).fill = tint
    s.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rows_for_export(d):
    lvl_label = {l["key"]: l["label"] for l in d["levels"]}
    out = {k: [] for k in SHEETS}
    for w in d["work_types"]:
        out["Work types"].append({
            **w, "applies_to": ", ".join(lvl_label.get(a, a) for a in (w.get("applies_to") or [])),
            "keywords": ", ".join(w.get("keywords") or []),
            "per_item": "Per item" if w.get("per_item") else "Grouped"})
        for s in w.get("steps") or []:
            out["Steps"].append({**s, "work_type": w["code"]})
        for r in w.get("rates") or []:
            out["Productivity"].append({**r, "work_type": w["code"],
                                        "driving": "Yes" if r.get("driving") else "No"})
        for m in w.get("materials") or []:
            out["Work-type materials"].append({**m, "work_type": w["code"]})
    out["Logic rules"] = [dict(r) for r in d.get("rules") or []]
    out["Resources"] = [{**r, "kind": "Equipment" if r["kind"] == "equipment" else "Labour"}
                        for r in d["resources"]]
    out["Materials"] = [{**m, "keywords": ", ".join(m.get("keywords") or [])}
                        for m in d["materials"]]
    for x in d.get("extras") or []:
        if x["kind"] == "eng":
            out["Engineering"].append(x)
        else:
            out["T&C"].append({**x, "level_key": lvl_label.get(x.get("level_key") or "", "Project")})
    out["Disciplines"] = [dict(x) for x in d["disciplines"]]
    return out


# ═══════════════════════════ الاستيراد ═══════════════════════════

def _norm_head(s):
    return re.sub(r"[^a-z0-9%]", "", str(s or "").lower().replace("*", ""))


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _code(v):
    return _txt(v).upper()


def _numv(v, errs, where, default=None):
    if v is None or _txt(v) == "":
        return default
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        errs.append({**where, "error": f"'{_txt(v)}' is not a number"})
        return None


def read_workbook(data):
    """
    يقرأ الملفّ إلى قواميس خام وأخطاء بصفوفها.

    يعيد ‎(sheets, errors)‎: ‎sheets‎ قاموسٌ بعنوان الورقة ← قائمة صفوف، ولا
    تظهر فيه إلا الأوراق الموجودة فعلًا في الملفّ. الصفّ الذي ينقصه حقل
    إلزاميّ يُسجَّل خطأً ولا يُقرأ؛ والباقي يُقرأ.
    """
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets, errors = {}, []
    names = {_norm_head(n): n for n in wb.sheetnames}
    for title, cols in SHEETS.items():
        real = names.get(_norm_head(title))
        if not real:
            continue
        ws = wb[real]
        it = ws.iter_rows(values_only=True)
        head = next(it, None)
        if not head:
            sheets[title] = []
            continue
        pos = {}
        wanted = {_norm_head(h): k for k, h, *_ in cols}
        for j, h in enumerate(head):
            k = wanted.get(_norm_head(h))
            if k and k not in pos:
                pos[k] = j
        missing = [h for k, h, req, *_ in cols if req and k not in pos]
        if missing:
            errors.append({"sheet": title, "row": 1,
                           "error": "Missing column: " + ", ".join(missing)})
            continue
        rows = []
        for i, r in enumerate(it, 2):
            if r is None or all(_txt(v) == "" for v in r):
                continue
            rec = {k: (r[j] if j < len(r) else None) for k, j in pos.items()}
            empty_req = [h for k, h, req, *_ in cols if req and _txt(rec.get(k)) == ""]
            if empty_req:
                errors.append({"sheet": title, "row": i,
                               "error": "Empty: " + ", ".join(empty_req)})
                continue
            rec["_row"] = i
            rows.append(rec)
        sheets[title] = rows
    wb.close()
    return sheets, errors


def _words(v):
    return [w.strip().lower() for w in re.split(r"[,\n;]", _txt(v)) if w.strip()]


def merge(current, sheets, errors, mode="merge"):
    """
    يدمج ما في الملفّ مع القواميس الحاليّة، فيعيد الجسم الذي يُحفظ به.

    يعيد ‎(body, diff, warnings)‎. ‎body‎ فيه المفاتيح التي مسّها الملفّ فقط،
    فما لم يمسّه لا يُعاد كتابته.
    """
    warn = []
    replace = mode == "replace"
    lv_by = {}
    for l in current["levels"]:
        lv_by[l["key"].lower()] = l["key"]
        lv_by[l["label"].lower()] = l["key"]
    csi_ok = {c["division"] for c in current["csi"]}

    def where(t, r):
        return {"sheet": t, "row": r.get("_row")}

    def rel(v, t, r):
        x = _code(v) or "FS"
        if x not in REL:
            errors.append({**where(t, r), "error": f"Link '{_txt(v)}' must be FS, SS, FF or SF"})
            return None
        return x

    body, diff = {}, {}

    def keyed(name, sheet, cur_list, key, build):
        """قاموسٌ بمفتاح واحد: يُحدَّث بالمفتاح ويُضاف الجديد."""
        if sheet not in sheets:
            return None
        incoming = {}
        for r in sheets[sheet]:
            x = build(r)
            if x is None:
                continue
            k = x[key]
            if k in incoming:
                warn.append({**where(sheet, r), "warning": f"{k} appears twice; the last row wins"})
            incoming[k] = x
        old = {x[key]: x for x in cur_list}
        if replace:
            out = list(incoming.values())
        else:
            out = [({**old[k], **incoming[k]} if k in incoming else old[k]) for k in old]
            out += [v for k, v in incoming.items() if k not in old]
        added = sum(1 for k in incoming if k not in old)
        changed = sum(1 for k in incoming if k in old and _differs(old[k], incoming[k]))
        removed = sum(1 for k in old if k not in incoming) if replace else 0
        diff[name] = {"added": added, "updated": changed, "removed": removed,
                      "rows": len(incoming)}
        return out

    # ── التخصّصات، الموارد، المواد ──
    def b_disc(r):
        csi = _txt(r.get("csi_division"))
        if csi and len(csi) == 1:
            csi = "0" + csi
        if csi and csi not in csi_ok:
            warn.append({**where("Disciplines", r), "warning": f"CSI division {csi} is not in MasterFormat"})
        return {"code": _code(r["code"]), "name": _txt(r.get("name")), "csi_division": csi}
    v = keyed("disciplines", "Disciplines", current["disciplines"], "code", b_disc)
    if v is not None:
        body["disciplines"] = v

    def b_res(r):
        k = _txt(r.get("kind")).lower()
        kind = "equipment" if k.startswith("equip") or k == "plant" else "labor"
        c = _numv(r.get("unit_cost"), errors, where("Resources", r), 0.0)
        return None if c is None else {"code": _code(r["code"]), "name": _txt(r.get("name")),
                                       "kind": kind, "unit_cost": c}
    v = keyed("resources", "Resources", current["resources"], "code", b_res)
    if v is not None:
        body["resources"] = v

    def b_mat(r):
        w = where("Materials", r)
        c = _numv(r.get("unit_cost"), errors, w, 0.0)
        lt = _numv(r.get("lead_time_days"), errors, w, 0.0)
        if c is None or lt is None:
            return None
        return {"code": _code(r["code"]), "name": _txt(r.get("name")), "unit": _txt(r.get("unit")),
                "unit_cost": c, "lead_time_days": lt, "keywords": _words(r.get("keywords"))}
    v = keyed("materials", "Materials", current["materials"], "code", b_mat)
    if v is not None:
        body["materials"] = v

    # ── أنواع الأعمال وأولادها ──
    wt_sheets = ("Work types", "Steps", "Productivity", "Work-type materials")
    if any(s in sheets for s in wt_sheets):
        def b_wt(r):
            w = where("Work types", r)
            days = _numv(r.get("days"), errors, w, 1.0)
            if days is None:
                return None
            ap = []
            for a in re.split(r"[,;]", _txt(r.get("applies_to"))):
                a = a.strip().lower()
                if not a:
                    continue
                if a in lv_by:
                    ap.append(lv_by[a])
                else:
                    warn.append({**w, "warning": f"Level '{a}' is not one of this system's levels"})
            csi = _txt(r.get("csi_division"))
            if csi and len(csi) == 1:
                csi = "0" + csi
            return {"code": _code(r["code"]), "name": _txt(r.get("name")),
                    "discipline": _code(r.get("discipline")), "csi_division": csi,
                    "unit": _txt(r.get("unit")), "days": days, "applies_to": ap,
                    "keywords": _words(r.get("keywords")),
                    "per_item": _txt(r.get("per_item")).lower().replace(" ", "")
                                in ("peritem", "item", "yes", "true")}
        cur_wt = [dict(w) for w in current["work_types"]]
        if "Work types" in sheets:
            wts = keyed("work_types", "Work types", cur_wt, "code", b_wt)
            # في الاستبدال يبقى لنوع العمل الباقي خطواته ومعدّلاته ما لم تأتِ
            # أوراقها في الملفّ؛ فاستبدال قائمة الأنواع لا يمحو تفاصيلها.
            old_by = {w["code"]: w for w in cur_wt}
            for w in wts:
                o = old_by.get(w["code"], {})
                for k in ("steps", "rates", "materials"):
                    if k not in w:
                        w[k] = list(o.get(k) or [])
        else:
            wts = cur_wt
        by = {w["code"]: w for w in wts}
        for w in wts:
            for k in ("steps", "rates", "materials"):
                w.setdefault(k, [])
        res_ok = {r["code"] for r in body.get("resources", current["resources"])}
        mat_ok = {m["code"] for m in body.get("materials", current["materials"])}

        def children(sheet, field, build, label):
            if sheet not in sheets:
                return
            got = {}
            for r in sheets[sheet]:
                wt = _code(r["work_type"])
                if wt not in by:
                    errors.append({**where(sheet, r), "error": f"Work type {wt} is not in the dictionary"})
                    continue
                x = build(r, by[wt])
                if x is not None:
                    got.setdefault(wt, []).append(x)
            n_rows = sum(len(v) for v in got.values())
            touched = set(got) if not replace else set(by)
            changed = 0
            for wt in touched:
                new = got.get(wt, [])
                if _rows_differ(by[wt].get(field) or [], new):
                    changed += 1
                by[wt][field] = new
            diff[label] = {"rows": n_rows, "work_types": len(got), "changed_for": changed}

        def b_step(r, w):
            wh = where("Steps", r)
            wgt = _numv(r.get("weight"), errors, wh, 0.0)
            lag = _numv(r.get("lag_days"), errors, wh, 0.0)
            rt = rel(r.get("rel_type"), "Steps", r)
            if wgt is None or lag is None or rt is None:
                return None
            return {"code": _code(r["code"]), "name": _txt(r.get("name")) or _code(r["code"]),
                    "weight": wgt, "rel_type": rt, "lag_days": lag}
        children("Steps", "steps", b_step, "steps")

        def b_rate(r, w):
            wh = where("Productivity", r)
            res = _code(r["resource"])
            if res not in res_ok:
                errors.append({**wh, "error": f"Resource {res} is not in the dictionary"})
                return None
            h = _numv(r.get("hours_per_unit"), errors, wh, 0.0)
            c = _numv(r.get("crew_count"), errors, wh, 1.0)
            if h is None or c is None:
                return None
            st = _code(r.get("step"))
            if st and st not in {s["code"] for s in w["steps"]}:
                warn.append({**wh, "warning": f"Step {st} is not a step of {w['code']}; the rate covers the whole work type"})
                st = ""
            dv = _txt(r.get("driving")).lower()
            return {"resource": res, "step": st, "hours_per_unit": h, "crew_count": c or 1,
                    "driving": dv in ("yes", "y", "true", "1", "x")}
        children("Productivity", "rates", b_rate, "productivity")

        def b_wm(r, w):
            wh = where("Work-type materials", r)
            m = _code(r["material"])
            if m not in mat_ok:
                errors.append({**wh, "error": f"Material {m} is not in the dictionary"})
                return None
            q = _numv(r.get("qty_per_unit"), errors, wh, 0.0)
            return None if q is None else {"material": m, "qty_per_unit": q}
        children("Work-type materials", "materials", b_wm, "work-type materials")

        disc_ok = {x["code"] for x in body.get("disciplines", current["disciplines"])}
        for w in wts:
            if w.get("discipline") and w["discipline"] not in disc_ok:
                warn.append({"sheet": "Work types", "row": None,
                             "warning": f"{w['code']}: discipline {w['discipline']} is not in the dictionary"})
            if w["steps"]:
                s = sum(float(x.get("weight") or 0) for x in w["steps"])
                if abs(s - 100) > 0.01:
                    warn.append({"sheet": "Steps", "row": None,
                                 "warning": f"{w['code']}: step weights add up to {round(s, 2)}%, not 100%"})
        body["work_types"] = wts

    # ── القواعد ──
    if "Logic rules" in sheets:
        wt_ok = {w["code"] for w in body.get("work_types", current["work_types"])}
        incoming = {}
        for r in sheets["Logic rules"]:
            wh = where("Logic rules", r)
            p, q = _code(r["pred_wt"]), _code(r["succ_wt"])
            sc = (_txt(r.get("scope")) or "same").lower()
            if sc not in SCOPES:
                errors.append({**wh, "error": f"Scope '{sc}' must be one of {', '.join(SCOPES)}"})
                continue
            rt = rel(r.get("rel_type"), "Logic rules", r)
            lag = _numv(r.get("lag_days"), errors, wh, 0.0)
            if rt is None or lag is None:
                continue
            for c in (p, q):
                if c not in wt_ok:
                    warn.append({**wh, "warning": f"Work type {c} is not in the dictionary"})
            incoming[(p, q, sc)] = {"pred_wt": p, "succ_wt": q, "scope": sc,
                                    "rel_type": rt, "lag_days": lag, "origin": "manual"}
        old = {(r["pred_wt"], r["succ_wt"], r["scope"]): r for r in current["rules"]}
        if replace:
            out = list(incoming.values())
        else:
            out = [incoming.get(k, v) for k, v in old.items()]
            out += [v for k, v in incoming.items() if k not in old]
        diff["logic rules"] = {
            "added": sum(1 for k in incoming if k not in old),
            "updated": sum(1 for k in incoming if k in old and _differs(old[k], incoming[k])),
            "removed": sum(1 for k in old if k not in incoming) if replace else 0,
            "rows": len(incoming)}
        body["rules"] = out

    # ── الهندسة و T&C ──
    for kind, sheet, key in (("eng", "Engineering", "engineering"), ("tc", "T&C", "tc")):
        if sheet not in sheets:
            continue
        wt_ok = {w["code"] for w in body.get("work_types", current["work_types"])}
        cur_x = [x for x in current["extras"] if x["kind"] == kind]
        incoming = {}
        for r in sheets[sheet]:
            wh = where(sheet, r)
            wt = _code(r["work_type"])
            if wt not in wt_ok:
                errors.append({**wh, "error": f"Work type {wt} is not in the dictionary"})
                continue
            d = _numv(r.get("duration_days"), errors, wh, 1.0)
            if d is None:
                continue
            x = {"work_type": wt, "code": _code(r["code"]), "name": _txt(r.get("name")),
                 "duration_days": d, "rel_type": "FS", "lag_days": 0.0, "level_key": ""}
            if kind == "eng":
                x["rel_type"] = rel(r.get("rel_type"), sheet, r)
                x["lag_days"] = _numv(r.get("lag_days"), errors, wh, 0.0)
                if x["rel_type"] is None or x["lag_days"] is None:
                    continue
            else:
                lv = _txt(r.get("level_key")).lower()
                if lv and lv not in ("project", "once", "once per project"):
                    if lv not in lv_by:
                        errors.append({**wh, "error": f"Level '{lv}' is not one of this system's levels"})
                        continue
                    x["level_key"] = lv_by[lv]
            incoming[(wt, x["code"])] = x
        old = {(x["work_type"], x["code"]): x for x in cur_x}
        if replace:
            out = list(incoming.values())
        else:
            out = [({**old[k], **incoming[k]} if k in incoming else old[k]) for k in old]
            out += [v for k, v in incoming.items() if k not in old]
        diff[key] = {"added": sum(1 for k in incoming if k not in old),
                     "updated": sum(1 for k in incoming if k in old and _differs(old[k], incoming[k])),
                     "removed": sum(1 for k in old if k not in incoming) if replace else 0,
                     "rows": len(incoming)}
        body[key] = out

    return body, diff, warn


def _rows_differ(old, new):
    """قائمتا صفوف لنوع عمل واحد: هل اختلفتا فيما يحمله الملفّ؟"""
    if len(old) != len(new):
        return True
    return any(_differs(o, n) for o, n in zip(old, new))


def _differs(a, b):
    """هل غيّر الملفّ شيئًا في الصفّ؟ تُقارَن الحقول التي يحملها الملفّ فقط."""
    for k, v in b.items():
        if k.startswith("_"):
            continue
        ov = a.get(k)
        if isinstance(v, float) or isinstance(ov, (int, float)):
            try:
                if abs(float(ov or 0) - float(v or 0)) > 1e-9:
                    return True
                continue
            except (TypeError, ValueError):
                pass
        if (ov or "") != (v or ""):
            return True
    return False
